"""Экспорт отчёта: Excel, CSV, XML, PDF (matplotlib) и PowerPoint (python-pptx, нативные диаграммы).

Содержимое собирается на странице: KPI-карточки, таблицы и графики Plotly (ui.collect_*).
Графики для PDF перерисовываются matplotlib по данным трасс Plotly; в PowerPoint — нативными
диаграммами, которые можно редактировать.
"""

from __future__ import annotations

import io
import zipfile
from datetime import datetime

import numpy as np
import pandas as pd

MIME = {
    "Excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "CSV": "text/csv",
    "ZIP": "application/zip",
    "XML": "application/xml",
    "PDF": "application/pdf",
    "PowerPoint": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def _slug(title: str) -> str:
    tr = str.maketrans("абвгдеёжзийклмнопрстуфхцчшщъыьэюя ", "abvgdeejziiklmnoprstufhccss_y_eua_")
    return "".join(ch for ch in title.lower().translate(tr) if ch.isalnum() or ch == "_")[:40] or "otchet"


def _filters_text(f) -> str:
    if f is None:
        return ""
    return f"Период: {f.start:%d.%m.%Y} — {f.end:%d.%m.%Y}"


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [str(c) if str(c) else "Статус" for c in out.columns]
    for c in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[c]) and getattr(out[c].dt, "tz", None):
            out[c] = out[c].dt.tz_localize(None)
    return out


def build(fmt: str, exp: dict) -> tuple[bytes, str, str]:
    title = exp.get("title", "Отчёт")
    base = f"{_slug(title)}_{datetime.now():%Y%m%d_%H%M}"
    tables = {k: _clean(v) for k, v in exp.get("tables", {}).items() if isinstance(v, pd.DataFrame)}
    kpis = exp.get("kpis", [])
    if kpis:
        tables = {"Показатели": pd.DataFrame(kpis, columns=["Показатель", "Значение", "Изменение"]), **tables}
    if fmt == "Excel":
        from bi.ui import to_excel
        return to_excel(tables or {"Пусто": pd.DataFrame()}), base + ".xlsx", MIME["Excel"]
    if fmt == "CSV":
        if len(tables) <= 1:
            df = next(iter(tables.values())) if tables else pd.DataFrame()
            return df.to_csv(index=False, sep=";").encode("utf-8-sig"), base + ".csv", MIME["CSV"]
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for i, (name, df) in enumerate(tables.items(), start=1):
                z.writestr(f"{i:02d}_{_slug(name)}.csv", df.to_csv(index=False, sep=";").encode("utf-8-sig"))
        return buf.getvalue(), base + "_csv.zip", MIME["ZIP"]
    if fmt == "XML":
        return _xml(title, tables), base + ".xml", MIME["XML"]
    if fmt == "PDF":
        return _pdf(title, exp, kpis, tables), base + ".pdf", MIME["PDF"]
    if fmt == "PowerPoint":
        return _pptx(title, exp, kpis, tables), base + ".pptx", MIME["PowerPoint"]
    raise ValueError(fmt)


# ---------------------------------------------------------------------------
def _xml(title: str, tables: dict[str, pd.DataFrame]) -> bytes:
    from xml.sax.saxutils import escape, quoteattr
    out = [f'<?xml version="1.0" encoding="UTF-8"?>\n<report title={quoteattr(title)} '
           f'generated="{datetime.now():%Y-%m-%dT%H:%M:%S}">']
    for name, df in tables.items():
        out.append(f"  <table name={quoteattr(name)}>")
        cols = list(df.columns)
        for row in df.itertuples(index=False):
            cells = "".join(f"<cell name={quoteattr(str(c))}>{escape('' if pd.isna(v) else str(v))}</cell>"
                            for c, v in zip(cols, row))
            out.append(f"    <row>{cells}</row>")
        out.append("  </table>")
    out.append("</report>")
    return "\n".join(out).encode("utf-8")


# ---------------------------------------------------------------------------
def _is_num(vals) -> bool:
    return pd.api.types.is_numeric_dtype(pd.Series(list(vals)).infer_objects())


def _traces(fig) -> list[dict]:
    """Простые трассы Plotly (bar/scatter) → категория (cat) и значение (val) для перерисовки.
    horizontal: значения по оси X, категории по оси Y (горизонтальные столбцы и точки к ним)."""
    res = []
    for tr in fig.data:
        t = tr.type
        if t not in ("bar", "scatter"):
            continue
        x = list(tr.x) if tr.x is not None else []
        y = list(tr.y) if tr.y is not None else []
        if not len(x) or not len(y) or len(x) != len(y):
            continue
        if getattr(tr, "fill", None) == "toself":
            continue
        horizontal = (getattr(tr, "orientation", None) == "h") or (_is_num(x) and not _is_num(y))
        cat, val = (y, x) if horizontal else (x, y)
        if not _is_num(val):
            continue
        res.append({"type": t, "name": tr.name or "", "cat": cat, "val": [None if v is None else float(v) for v in val],
                    "horizontal": horizontal, "yaxis": getattr(tr, "yaxis", None) or "y"})
    return res


def _fig_title(fig, default: str) -> str:
    try:
        return fig.layout.title.text or default
    except Exception:
        return default


def _pdf(title: str, exp: dict, kpis: list, tables: dict) -> bytes:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "axes.spines.top": False,
                         "axes.spines.right": False})
    colors = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#6250d6", "#e34948"]
    buf = io.BytesIO()
    with PdfPages(buf) as pdf:
        # Титульная страница с KPI
        fig = plt.figure(figsize=(11.69, 8.27))
        fig.text(0.05, 0.92, title, fontsize=20, weight="bold")
        fig.text(0.05, 0.875, f"{_filters_text(exp.get('filters'))} · сформировано {datetime.now():%d.%m.%Y %H:%M} · "
                 "данные синтетические", fontsize=10, color="#52514e")
        for i, (label, value, delta) in enumerate(kpis[:15]):
            r, c = divmod(i, 3)
            x, y = 0.05 + c * 0.31, 0.78 - r * 0.13
            fig.patches.append(plt.Rectangle((x - 0.01, y - 0.075), 0.29, 0.11, transform=fig.transFigure,
                                             facecolor="#f6f9fd", edgecolor="#e1e0d9"))
            fig.text(x, y + 0.012, label, fontsize=9, color="#52514e")
            fig.text(x, y - 0.035, value, fontsize=16, weight="bold")
            if delta:
                fig.text(x + 0.15, y - 0.035, delta, fontsize=9, color="#52514e")
        if not kpis:
            fig.text(0.05, 0.75, "На странице нет KPI-карточек.", fontsize=11)
        pdf.savefig(fig)
        plt.close(fig)
        # Графики
        for name, f in exp.get("figures", {}).items():
            tr = _traces(f)
            if not tr:
                continue
            fig, ax = plt.subplots(figsize=(11.69, 8.27))
            horizontal = any(t["horizontal"] for t in tr)
            tr = [t for t in tr if t["horizontal"] == horizontal][:6]
            # единый тип оси категорий для всех рядов: даты или текст (иначе matplotlib падает на смеси)
            xkind = "cat"
            if not horizontal and not any(_is_num(t["cat"]) for t in tr):
                try:
                    conv = [list(pd.to_datetime(pd.Series(t["cat"]), errors="raise")) for t in tr]
                    for t, xv in zip(tr, conv):
                        t["cat"] = xv
                    xkind = "date"
                except (ValueError, TypeError, OverflowError):
                    pass
            elif not horizontal and all(_is_num(t["cat"]) for t in tr):
                xkind = "num"
            ax0, ax2 = ax, None
            for k, t in enumerate(tr):
                col = colors[k % len(colors)]
                y = [np.nan if v is None else v for v in t["val"]]
                # ряд со своей осью Y (вторая шкала / подграфик) — на правую ось
                if not horizontal and t["yaxis"] != tr[0]["yaxis"]:
                    ax2 = ax2 or ax0.twinx()
                    ax = ax2
                else:
                    ax = ax0
                xv = t["cat"] if xkind in ("date", "num") else [str(v)[:24] for v in t["cat"]]
                if t["type"] == "bar":
                    if horizontal:
                        ax.barh(xv, y, color=col, label=t["name"], alpha=0.9 if k == 0 else 0.6)
                    elif xkind == "date":
                        ax.bar(xv, y, color=col, label=t["name"], width=0.8 if len(xv) > 60 else 5)
                    else:
                        ax.bar(xv, y, color=col, label=t["name"])
                elif horizontal:
                    ax.plot(y, xv, color=col, linestyle="none", marker="D", label=t["name"])
                else:
                    ax.plot(xv, y, color=col, linewidth=2, label=t["name"])
            ax = ax0
            ax.set_title(_fig_title(f, name), loc="left", fontsize=13, weight="bold")
            ax.grid(axis="x" if horizontal else "y", color="#e1e0d9")
            if len(tr) > 1:
                h1, l1 = ax0.get_legend_handles_labels()
                h2, l2 = ax2.get_legend_handles_labels() if ax2 is not None else ([], [])
                ax.legend(h1 + h2, l1 + l2, frameon=False)
            if not horizontal:
                plt.setp(ax.get_xticklabels(), rotation=30, ha="right")
            fig.tight_layout()
            pdf.savefig(fig)
            plt.close(fig)
        # Таблицы (первые 30 строк каждой)
        for name, df in tables.items():
            if name == "Показатели" or df.empty:
                continue
            d = df.head(30).copy()
            d = d.iloc[:, :9]
            for c in d.columns:
                if pd.api.types.is_float_dtype(d[c]):
                    d[c] = d[c].map(lambda v: "" if pd.isna(v) else (f"{v:.3f}" if abs(v) < 10 else f"{v:,.0f}"
                                                                    .replace(",", " ")))
            d = d.astype(str).apply(lambda s: s.str.slice(0, 28))
            fig, ax = plt.subplots(figsize=(11.69, 8.27))
            ax.axis("off")
            ax.set_title(f"{name}" + (" (первые 30 строк)" if len(df) > 30 else ""), loc="left", fontsize=13,
                         weight="bold")
            tbl = ax.table(cellText=d.values, colLabels=[str(c)[:22] for c in d.columns], loc="upper left",
                           cellLoc="left", colLoc="left")
            tbl.auto_set_font_size(False)
            tbl.set_fontsize(7.5)
            tbl.scale(1, 1.25)
            for (r, _), cell in tbl.get_celld().items():
                cell.set_edgecolor("#e1e0d9")
                if r == 0:
                    cell.set_facecolor("#eef5fd")
                    cell.set_text_props(weight="bold")
            pdf.savefig(fig)
            plt.close(fig)
    return buf.getvalue()


# ---------------------------------------------------------------------------
def _pptx(title: str, exp: dict, kpis: list, tables: dict) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    blue = RGBColor(0x2A, 0x78, 0xD6)

    def add_title(slide, text, sub=None):
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8))
        p = tb.text_frame.paragraphs[0]
        p.text = text
        p.font.size, p.font.bold = Pt(26), True
        if sub:
            p2 = tb.text_frame.add_paragraph()
            p2.text = sub
            p2.font.size = Pt(12)
            p2.font.color.rgb = RGBColor(0x52, 0x51, 0x4E)

    s = prs.slides.add_slide(blank)
    add_title(s, title, f"{_filters_text(exp.get('filters'))} · сформировано {datetime.now():%d.%m.%Y %H:%M} · "
                        "данные синтетические")
    for i, (label, value, delta) in enumerate(kpis[:12]):
        r, c = divmod(i, 4)
        box = s.shapes.add_textbox(Inches(0.5 + c * 3.1), Inches(1.6 + r * 1.75), Inches(2.9), Inches(1.5))
        tf = box.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = label
        tf.paragraphs[0].font.size = Pt(12)
        p = tf.add_paragraph()
        p.text = value
        p.font.size, p.font.bold = Pt(24), True
        p.font.color.rgb = blue
        if delta:
            p3 = tf.add_paragraph()
            p3.text = delta
            p3.font.size = Pt(11)

    for name, f in exp.get("figures", {}).items():
        tr = _traces(f)
        if not tr:
            continue
        horizontal = tr[0]["horizontal"]
        tr = [t for t in tr if t["horizontal"] == horizontal]
        cats = [str(v)[:24] if not isinstance(v, (pd.Timestamp, np.datetime64)) else pd.Timestamp(v).strftime("%d.%m.%y")
                for v in tr[0]["cat"]]
        cd = CategoryChartData()
        cd.categories = cats[:60]
        kind = None
        n_series = 0
        for t in tr[:5]:
            if len(t["val"]) != len(cats):
                continue
            cd.add_series(t["name"] or name, [v if v is not None and np.isfinite(v) else None for v in t["val"][:60]])
            kind = kind or t["type"]
            n_series += 1
        if kind is None:
            continue
        s = prs.slides.add_slide(blank)
        add_title(s, _fig_title(f, name))
        ctype = (XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED) \
            if kind == "bar" else XL_CHART_TYPE.LINE
        gframe = s.shapes.add_chart(ctype, Inches(0.5), Inches(1.3), Inches(12.3), Inches(5.8), cd)
        ch = gframe.chart
        ch.has_legend = n_series > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
        ch.font.size = Pt(10)

    for name, df in tables.items():
        if name == "Показатели" or df.empty:
            continue
        d = df.head(15).iloc[:, :8]
        s = prs.slides.add_slide(blank)
        add_title(s, name, "первые 15 строк; полная таблица — в выгрузке Excel" if len(df) > 15 else None)
        rows, cols = len(d) + 1, len(d.columns)
        tbl = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.4), Inches(12.3), Inches(0.35 * rows)).table
        for j, c in enumerate(d.columns):
            tbl.cell(0, j).text = str(c)[:30]
        for i, row in enumerate(d.itertuples(index=False), start=1):
            for j, v in enumerate(row):
                if isinstance(v, float):
                    txt = "" if pd.isna(v) else (f"{v:.3f}" if abs(v) < 10 else f"{v:,.0f}".replace(",", " "))
                else:
                    txt = str(v)
                tbl.cell(i, j).text = txt[:40]
        for i in range(rows):
            for j in range(cols):
                for p in tbl.cell(i, j).text_frame.paragraphs:
                    p.font.size = Pt(10)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
